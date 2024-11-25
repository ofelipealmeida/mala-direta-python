import pandas as pd
from docx import Document
import os

# Nome do arquivo CSV
arquivo_csv = "bbb.csv"

# Nome do modelo de documento
arquivo_modelo = "modelo.docx"

# Caminho da pasta onde os documentos gerados serão salvos
pasta_saida = "/media/felipe/Novo volume/Ubuntu/Projects/Gerador de Assinaturas/Documentos"
os.makedirs(pasta_saida, exist_ok=True)  # Garante que a pasta exista

# Carregar a base de dados
df = pd.read_csv(arquivo_csv, sep=";")  # Ler CSV separado por ponto e vírgula

# Gerar documentos personalizados
for index, linha in df.iterrows():
    # Criar uma cópia do modelo
    doc = Document(arquivo_modelo)
    
    # Substituir os placeholders no texto do documento
    for paragrafo in doc.paragraphs:
        paragrafo.text = paragrafo.text.replace("{{{nomes}}}", str(linha["nomes"]))
        paragrafo.text = paragrafo.text.replace("{{{cpf}}}", str(linha["cpf"]))
        paragrafo.text = paragrafo.text.replace("{{{valores}}}", str(linha["valores"]))
        paragrafo.text = paragrafo.text.replace("{{{data}}}", str(linha["data"]))

    # Salvar o documento gerado com o nome no formato "nome_id.docx"
    nome_arquivo = f"{linha['nomes']}_{linha['id']}.docx"
    caminho_saida = os.path.join(pasta_saida, nome_arquivo)
    doc.save(caminho_saida)

print(f"Documentos gerados na pasta: {pasta_saida}")
